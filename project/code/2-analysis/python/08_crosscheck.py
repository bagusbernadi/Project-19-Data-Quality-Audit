"""08: Cek silang. (a) audit R01-R06 pandas vs loop Python murni (himpunan entri per aturan), (b) vs SQL DuckDB (tambahan, rancangan tidak memakai SQL),
(c) rekap mentah vs Excel (SUMIFS, dihitung ulang LibreOffice). R07 dicek silang lewat loop rantai terpisah.
"""
import re
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import duckdb
import numpy as np
import openpyxl
import pandas as pd
import audit_lib as AL
import common as CM

_h = pd.read_csv(CM.RES1 / "log_entries.csv", nrows=0).columns
rc = [c for c in _h if c.endswith("_raw")]
raw = pd.read_csv(CM.RES1 / "log_entries.csv", dtype={c: str for c in rc})
raw[rc] = raw[rc].fillna("")
al = pd.read_csv(CM.RES2 / "anomaly_log.csv")
out = []


def isint(s):
    try:
        v = float(s)
    except (TypeError, ValueError):
        return False
    return v >= 0 and v == int(v) and bool(re.fullmatch(r"-?\d+(\.0+)?", str(s).strip()))


def num(v):
    return None if v is None or (isinstance(v, float) and np.isnan(v)) else v


# (a) loop Python murni
sets = {}
seen = set()
for r in raw.itertuples(index=False):
    d = r._asdict(); ver = d["version"]; eid = d["entry_id"]
    def n(f):  # nilai numerik valid atau None
        x = d[f]; return x if (isinstance(x, (int, float)) and not np.isnan(x) and x >= 0 and x == int(x)) else None
    tipe = {f: d[f + "_tipe"] for f in AL.NUMF}
    awal = n("sisa_awal") or 0; ok = n("qty_ok") or 0; sc = n("qty_scrap") or 0; rcv = n("qty_received")
    if rcv is not None and ok + sc > rcv + awal:
        sets.setdefault((ver, "R01"), set()).add(eid)
    for f in AL.NUMF:
        if tipe[f] in ("int", "float", "teks") and n(f) is None:
            sets.setdefault((ver, "R02"), set()).add(eid)
    if any(d[f + "_raw"] == "" for f in AL.REQ):
        sets.setdefault((ver, "R03"), set()).add(eid)
    for f, items in (("shift", CM.SHIFTS), ("model", CM.MODELS), ("ng_origin", CM.STATIONS)):
        if d[f + "_raw"] != "" and d[f + "_raw"] not in items:
            sets.setdefault((ver, "R04"), set()).add(eid)
    if d["sisa_is_formula"] == 0:
        sets.setdefault((ver, "R05"), set()).add(eid)
    k = (ver, d["date_raw"], d["shift_raw"], d["model_raw"], d["ng_origin_raw"], d["batch_raw"])
    if k in seen:
        sets.setdefault((ver, "R06"), set()).add(eid)
    seen.add(k)
for ver in ("before", "after"):
    for rule in ("R01", "R02", "R03", "R04", "R05", "R06"):
        a = set(al[(al.version == ver) & (al.rule_id == rule)].entry_id)
        b = sets.get((ver, rule), set())
        out.append(dict(cek="pandas vs loop Python", version=ver, item=rule, nilai_a=len(a), nilai_b=len(b), cocok=int(a == b)))

# R07 loop rantai
for ver in ("before", "after"):
    d = raw[raw.version == ver].copy(); d["m"], d["s"] = AL.norm_pair(d)
    d = d[d.m.notna() & d.s.notna()]
    bad = 0
    for (m, s), g in d.groupby(["m", "s"]):
        last = {}
        for mon in sorted(g.month.unique()):
            gm = g[g.month == mon].sort_values("excel_row")
            f = gm.iloc[0]
            if last:
                prev = last["v"]
                if prev is not None and not np.isnan(f.sisa_awal) and f.sisa_awal != prev:
                    bad += 1
            l = gm.iloc[-1]
            vals = [l.sisa_awal, l.qty_received, l.qty_ok, l.qty_scrap]
            last = {"v": None if any(np.isnan(v) or v < 0 or v != int(v) for v in vals) else vals[0] + vals[1] - vals[2] - vals[3]}
    a = int(((al.version == ver) & (al.rule_id == "R07")).sum())
    out.append(dict(cek="pandas vs loop Python", version=ver, item="R07", nilai_a=a, nilai_b=bad, cocok=int(a == bad)))

# (b) SQL DuckDB
con = duckdb.connect()
con.register("log", raw)
q = lambda w: con.sql(w).df()
def cnt(rule_sql):
    return {(r.version): int(r.n) for r in q(f"select version, count(distinct entry_id) n from ({rule_sql}) group by version").itertuples()}
num_ok = lambda c: f"(try_cast(nullif({c},'') as double) is not null and try_cast({c} as double) >= 0 and try_cast({c} as double) = floor(try_cast({c} as double)))"
sql = {
    "R03": "select version, entry_id from log where " + " or ".join(f"{f}_raw=''" for f in AL.REQ),
    "R04": "select version, entry_id from log where (shift_raw<>'' and shift_raw not in ('Shift 1','Shift 2')) or (model_raw<>'' and model_raw not in ('PCB-A1','PCB-A2','PCB-B1','PCB-B2')) or (ng_origin_raw<>'' and ng_origin_raw not in ('Stasiun SMT','Stasiun AOI','Stasiun ICT'))",
    "R05": "select version, entry_id from log where sisa_is_formula=0",
    "R06": "select version, entry_id from (select version, entry_id, row_number() over (partition by version,date_raw,shift_raw,model_raw,ng_origin_raw,batch_raw order by excel_row) rn from log) where rn>1",
    "R01": "select version, entry_id from log where qty_received is not null and coalesce(case when qty_ok>=0 and qty_ok=floor(qty_ok) then qty_ok end,0)+coalesce(case when qty_scrap>=0 and qty_scrap=floor(qty_scrap) then qty_scrap end,0) > qty_received+coalesce(case when sisa_awal>=0 and sisa_awal=floor(sisa_awal) then sisa_awal end,0)",
}
for rule, s in sql.items():
    c = cnt(s)
    for ver in ("before", "after"):
        a = int(((al.version == ver) & (al.rule_id == rule)).sum()) if rule != "R03" else int(al[(al.version == ver) & (al.rule_id == rule)].entry_id.nunique())
        a = al[(al.version == ver) & (al.rule_id == rule)].entry_id.nunique()
        out.append(dict(cek="pandas vs SQL DuckDB", version=ver, item=rule, nilai_a=int(a), nilai_b=c.get(ver, 0), cocok=int(a == c.get(ver, 0))))

# (c) Excel: rekap mentah (SUMIFS) dihitung ulang LibreOffice
rec = pd.read_csv(CM.RES2 / "recap_reconcile.csv")
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "data"
ws.append(["version", "month", "diterima", "ok", "scrap", "sisa"])
for r in raw.itertuples():
    ws.append([r.version, r.month, *(float(getattr(r, f)) if getattr(r, f + "_tipe") in ("int", "float") and not np.isnan(getattr(r, f)) else "" for f in ("qty_received", "qty_ok", "qty_scrap")),
               float(pd.to_numeric(r.sisa, errors="coerce")) if not np.isnan(pd.to_numeric(r.sisa, errors="coerce")) else ""])
wr = wb.create_sheet("rekap"); wr.append(["version", "month", "diterima", "ok", "scrap", "sisa"])
for i, r in enumerate(rec.itertuples(), start=2):
    wr.append([r.version, r.bulan] + [f'=SUMIFS(data!{c}:{c},data!$A:$A,$A{i},data!$B:$B,$B{i})' for c in "CDEF"])
tmp = Path(tempfile.mkdtemp()); src = tmp / "x.xlsx"; wb.save(src); (tmp / "o").mkdir()
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp / "o"), str(src)], check=True, capture_output=True)
v = openpyxl.load_workbook(tmp / "o" / "x.xlsx", data_only=True)["rekap"]
for i, r in enumerate(rec.itertuples(), start=2):
    for c, f in zip("CDEF", ("diterima", "ok", "scrap", "sisa")):
        a = float(getattr(r, f + "_mentah")); b = float(v[f"{c}{i}"].value or 0)
        out.append(dict(cek="pandas vs Excel (LibreOffice)", version=r.version, item=f"{r.bulan} {f}", nilai_a=a, nilai_b=b, cocok=int(abs(a - b) < 1e-6)))
res = pd.DataFrame(out)
res.to_csv(CM.RES2 / "crosscheck.csv", index=False)
print(res.groupby("cek").cocok.agg(["sum", "size"]))
assert res.cocok.all(), res[res.cocok == 0]
