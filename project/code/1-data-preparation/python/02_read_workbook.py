"""02: Baca tiap workbook dua kali dengan openpyxl: sebagai rumus (data_only=False) dan sebagai nilai terhitung (data_only=True).

openpyxl menulis rumus tanpa nilai tersimpan, jadi workbook dihitung ulang oleh LibreOffice (headless) ke salinan sementara,
lalu nilainya dibaca. Hasil (data/clean, tidak di-commit): cell_raw.csv (satu baris per sel: nilai, tipe) dan
results/1-data-preparation/formula_audit.csv (satu baris per sel kolom Sisa: tabel formula_audit di rancangan).
"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import numpy as np
import openpyxl
import pandas as pd
import common as CM


def kind(v):
    if v is None:
        return "kosong"
    if isinstance(v, bool):
        return "bool"
    if isinstance(v, int):
        return "int"
    if isinstance(v, float):
        return "float"
    if hasattr(v, "year"):
        return "tanggal"
    s = str(v)
    return "rumus" if s.startswith("=") else "teks"


def recalc(src, outdir):
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(outdir), str(src)],
                   check=True, capture_output=True, timeout=240)
    return outdir / Path(src).name


tmp = Path(tempfile.mkdtemp())
cells, fa = [], []
for name, ver in CM.WORKBOOKS.items():
    src = CM.SIM / f"{name}.xlsx"
    wf = openpyxl.load_workbook(src)["Log"]
    wv = openpyxl.load_workbook(recalc(src, tmp), data_only=True)["Log"]
    head = [c.value for c in wf[1]]
    assert head == CM.HEADER, head
    for r in range(2, wf.max_row + 1):
        for j, h in enumerate(CM.HEADER, start=1):
            f, v = wf.cell(r, j).value, wv.cell(r, j).value
            raw = f if kind(f) == "rumus" else v
            cells.append(dict(version=ver, workbook=name, excel_row=r, column=h, value="" if raw is None else (raw.isoformat() if hasattr(raw, "isoformat") else str(raw)),
                              tipe=kind(raw), nilai_terhitung="" if v is None else (v.isoformat() if hasattr(v, "isoformat") else str(v))))
        # audit rumus kolom Sisa
        f = wf.cell(r, CM.HEADER.index("Sisa") + 1).value
        v = wv.cell(r, CM.HEADER.index("Sisa") + 1).value
        nums = [wf.cell(r, CM.HEADER.index(c) + 1).value for c in ("Sisa Awal", "Qty Diterima", "Qty OK", "Qty Scrap")]
        ok = all(isinstance(x, (int, float)) and not isinstance(x, bool) for x in nums)
        rec = nums[0] + nums[1] - nums[2] - nums[3] if ok else np.nan
        fa.append(dict(version=ver, sheet="Log", cell=f"J{r}", column="Sisa", expected_formula=f"=F{r}+G{r}-H{r}-I{r}", is_formula=int(kind(f) == "rumus"),
                       formula_found=str(f) if kind(f) == "rumus" else "", value=v, recomputed_value=rec))
cr = pd.DataFrame(cells)
cr.to_csv(CM.CLEAN / "cell_raw.csv", index=False)
fa = pd.DataFrame(fa)
fa["value"] = pd.to_numeric(fa.value, errors="coerce")
fa.to_csv(CM.RES1 / "formula_audit.csv", index=False, float_format="%.10g")
shutil.rmtree(tmp, ignore_errors=True)
print(cr.groupby("version").excel_row.nunique().to_dict(), "baris;", len(fa), "sel Sisa;", int(fa.is_formula.sum()), "berisi rumus")
