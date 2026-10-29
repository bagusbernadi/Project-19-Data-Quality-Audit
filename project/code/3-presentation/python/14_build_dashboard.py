"""14: Dashboard Excel (Audit_Dashboard.xlsx, grafik native openpyxl) dan PNG pareto + before-after. Semua angka dari CSV hasil."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
import common as CM

R2 = CM.RES2
rc = pd.read_csv(R2 / "audit_rule_counts.csv"); kp = pd.read_csv(R2 / "kpi.csv"); ba = pd.read_csv(R2 / "before_after.csv").iloc[0]
col = pd.read_csv(R2 / "audit_column.csv"); rm = pd.read_csv(R2 / "recap_reconcile.csv"); hs = pd.read_csv(R2 / "hours_saved.csv").iloc[0]
wb = Workbook(); ws = wb.active; ws.title = "Ringkasan"
H = PatternFill("solid", fgColor="1A73E8")
def table(ws, r0, df, title):
    ws.cell(r0, 1, title).font = Font(bold=True, size=12)
    for j, c in enumerate(df.columns, 1):
        x = ws.cell(r0 + 1, j, c); x.font = Font(bold=True, color="FFFFFF"); x.fill = H
    for i, row in enumerate(df.itertuples(index=False), r0 + 2):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else v)
    return r0 + 1, r0 + 1 + len(df)
ws["A1"] = "Audit Dashboard: log rework (data simulasi)"; ws["A1"].font = Font(bold=True, size=14)
k = pd.DataFrame([["Error per 1.000 entri, log lama", ba.error_per_1000_before], ["Error per 1.000 entri, template", ba.error_per_1000_after],
                  ["Penurunan", ba.penurunan], ["Jam koreksi hemat per bulan (estimated)", hs.jam_hemat_per_bulan]], columns=["Indikator", "Nilai"])
table(ws, 3, k, "Indikator utama")
r_before = rc[rc.version == "before"][["rule_id", "pelanggaran"]].sort_values("pelanggaran", ascending=False)
a, b = table(ws, 10, r_before, "Pelanggaran per aturan (log lama)")
ch = BarChart(); ch.title = "Pelanggaran per aturan, log lama"; ch.add_data(Reference(ws, min_col=2, min_row=a, max_row=b), titles_from_data=True); ch.set_categories(Reference(ws, min_col=1, min_row=a + 1, max_row=b)); ch.height, ch.width = 7, 14
ws.add_chart(ch, "E3")
cb = col[col.version == "before"][["column", "pelanggaran"]]
a, b = table(ws, 22, cb, "Pelanggaran per kolom (log lama)")
ch2 = BarChart(); ch2.type = "bar"; ch2.title = "Pelanggaran per kolom"; ch2.add_data(Reference(ws, min_col=2, min_row=a, max_row=b), titles_from_data=True); ch2.set_categories(Reference(ws, min_col=1, min_row=a + 1, max_row=b)); ch2.height, ch2.width = 7, 14
ws.add_chart(ch2, "E18")
kk = kp[["version", "salah_isi_per_1000", "dilewati_per_1000", "bypass_per_1000"]]
a, b = table(ws, 36, kk, "Error per 1.000 entri menurut kategori")
ch3 = BarChart(); ch3.title = "Error per 1.000 entri: before vs after"; ch3.add_data(Reference(ws, min_col=2, max_col=4, min_row=a, max_row=b), titles_from_data=True); ch3.set_categories(Reference(ws, min_col=1, min_row=a + 1, max_row=b)); ch3.height, ch3.width = 7, 14
ws.add_chart(ch3, "E33")
ws2 = wb.create_sheet("Rekonsiliasi"); table(ws2, 1, rm[["version", "bulan", "ok_mentah", "ok_koreksi", "scrap_mentah", "scrap_koreksi", "sisa_mentah", "sisa_koreksi"]], "Rekap mentah vs terkoreksi")
for w in (ws, ws2):
    w.column_dimensions["A"].width = 42
    for c in "BCDEFGH": w.column_dimensions[c].width = 16
fl = pd.read_csv(R2 / "entry_flags.csv"); fl["minggu"] = pd.to_datetime(fl.date, errors="coerce").dt.strftime("%G-W%V")
wk = fl.dropna(subset=["minggu"]).groupby(["version", "minggu"]).agg(entri=("entry_id", "size"), pelanggaran=("n_pelanggaran", "sum")).reset_index()
wk["error_per_1000"] = wk.pelanggaran / wk.entri * 1000
wk.to_csv(CM.RES3 / "weekly_trend.csv", index=False, float_format="%.10g")
from openpyxl.chart import LineChart
ws3 = wb.create_sheet("Tren mingguan")
pv = wk.pivot(index="minggu", columns="version", values="error_per_1000").reset_index().fillna(0)[["minggu", "before", "after"]]
a, b = table(ws3, 1, pv, "Error per 1.000 entri per minggu")
lc = LineChart(); lc.title = "Tren mingguan error per 1.000 entri"; lc.add_data(Reference(ws3, min_col=2, max_col=3, min_row=a, max_row=b), titles_from_data=True)
lc.set_categories(Reference(ws3, min_col=1, min_row=a + 1, max_row=b)); lc.height, lc.width = 8, 18; ws3.add_chart(lc, "E3")
wb.save(CM.OUTPUTS / "Audit_Dashboard.xlsx")

fig, axs = plt.subplots(1, 2, figsize=(11, 4.2))
d = r_before; axs[0].bar(d.rule_id, d.pelanggaran, color="#1A73E8"); axs[0].set_title("Pelanggaran per aturan (log lama)"); axs[0].set_ylabel("pelanggaran")
cats = ["salah_isi_per_1000", "dilewati_per_1000", "bypass_per_1000"]; w = 0.35
for i, (v, c) in enumerate((("before", "#EA4335"), ("after", "#34A853"))):
    axs[1].bar([x + i * w for x in range(3)], kp[kp.version == v][cats].iloc[0].values, w, label="log lama" if v == "before" else "template", color=c)
axs[1].set_xticks([x + w / 2 for x in range(3)]); axs[1].set_xticklabels(["salah isi", "dilewati", "bypass"]); axs[1].set_title("Error per 1.000 entri"); axs[1].legend()
fig.suptitle("Audit kualitas data log rework (data simulasi)"); fig.tight_layout(); fig.savefig(CM.OUTPUTS / "audit_summary.png", dpi=130)
print("dashboard ditulis")
